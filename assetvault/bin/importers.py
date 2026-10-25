"""CSV and XLSX import. Detects Intune, Entra, AD, and Jamf exports, or a generic sheet.

The detected mapping is returned to the caller. Nothing is written to the vault here.
"""
import csv, io, os

import avcore as av
from connectors.base import DEVICE_FIELDS, device
from errors import OpsError

_LABELS = {
    "intune": "Intune export",
    "entra": "Entra export",
    "ad": "AD export",
    "jamf": "Jamf export",
    "generic": "CSV import",
}
_MANAGED = {"intune": "Intune", "entra": "Entra ID", "ad": "Active Directory", "jamf": "Jamf"}

_INTUNE = {
    "name": ["device name"],
    "serial": ["serial number"],
    "mac": ["wi-fi mac", "wi-fi mac address", "wifi mac", "ethernet mac", "mac address"],
    "brand": ["manufacturer"],
    "model": ["model"],
    "os": ["os"],
    "os_version": ["os version"],
    "assigned_user": ["primary user upn", "primary user", "user principal name"],
    "last_checkin": ["last check-in", "last check-in time"],
    "encryption_status": ["encrypted"],
    "managed_by": ["managed by"],
    "external_id": ["azure ad device id", "intune device id", "device id"],
    "asset_tag": ["asset tag"],
}
_ENTRA = {
    "name": ["display name"],
    "os": ["os", "operating system"],
    "os_version": ["version", "os version"],
    "assigned_user": ["owner", "user", "registered owners"],
    "last_checkin": ["activity", "last activity"],
    "managed_by": ["mdm"],
    "external_id": ["device id"],
}
_AD = {
    "name": ["name", "dnshostname", "dns host name"],
    "os": ["operatingsystem", "operating system"],
    "os_version": ["operatingsystemversion", "operating system version"],
    "last_checkin": ["lastlogontimestamp", "last logon timestamp", "lastlogondate", "last logon date"],
    "notes": ["description"],
    "external_id": ["objectguid", "object guid"],
}
_JAMF = {
    "name": ["computer name", "device name"],
    "serial": ["serial number"],
    "mac": ["mac address", "wifi mac address", "wi-fi mac address"],
    "model": ["model", "model identifier"],
    "brand": ["make", "manufacturer"],
    "os": ["operating system", "os"],
    "os_version": ["operating system version", "os version"],
    "assigned_user": ["username", "full name", "email address", "email"],
    "last_checkin": ["last check-in", "last inventory update", "last contact time"],
    "encryption_status": ["filevault 2 status", "filevault status", "disk encryption"],
    "asset_tag": ["asset tag"],
    "external_id": ["udid", "jss computer id"],
}
_GENERIC = {
    "name": ["name", "device name", "computer name", "display name", "hostname", "host name"],
    "serial": ["serial", "serial number", "serialnumber"],
    "mac": ["mac", "mac address", "wifi mac", "wi-fi mac address"],
    "asset_tag": ["asset tag", "asset tag number"],
    "brand": ["brand", "manufacturer", "make"],
    "model": ["model", "model name"],
    "os": ["os", "operating system"],
    "os_version": ["os version", "operating system version"],
    "assigned_user": ["assigned user", "user", "username", "owner", "primary user"],
    "last_checkin": ["last check-in", "last checkin", "last seen"],
    "encryption_status": ["encryption", "encryption status", "encrypted"],
    "managed_by": ["managed by"],
    "external_id": ["external id", "external_id", "id", "device id"],
    "notes": ["notes", "description"],
}


def norm_header(value):
    return " ".join(str(value or "").replace("_", " ").split()).casefold()


def index_headers(headers):
    out = {}
    for header in headers:
        if not header:
            continue
        out.setdefault(norm_header(header), header)
    return out


def _has(index, *aliases):
    return any(norm_header(alias) in index for alias in aliases)


def _pick(index, aliases):
    for alias in aliases:
        key = norm_header(alias)
        if key in index:
            return index[key]
    return None


def _mapping(index, spec):
    out = {}
    for field, aliases in spec.items():
        header = _pick(index, aliases)
        if header:
            out[field] = header
    return out


def detect(headers):
    """Return (profile, mapping). profile is intune, entra, ad, jamf, or generic."""
    index = index_headers(headers)
    if _has(index, "device name") and _has(index, "managed by", "primary user upn", "compliance", "enrollment date"):
        return "intune", _mapping(index, _INTUNE)
    if _has(index, "display name") and _has(index, "device id") and _has(index, "join type", "mdm"):
        return "entra", _mapping(index, _ENTRA)
    if _has(index, "objectguid", "object guid") and _has(index, "operatingsystem", "operating system", "dnshostname", "dns host name", "distinguishedname", "distinguished name", "lastlogontimestamp", "last logon timestamp"):
        return "ad", _mapping(index, _AD)
    if _has(index, "serial number") and _has(index, "asset tag", "udid", "filevault 2 status", "filevault status", "jss computer id", "last inventory update", "model identifier"):
        return "jamf", _mapping(index, _JAMF)
    return "generic", _mapping(index, _GENERIC)


def read_table(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        with av.private_temp():
            return _xlsx(path)
    if ext == ".xls":
        raise OpsError("legacy .xls is not supported; save as .xlsx or .csv")
    return _csv(path)


def _csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read()
    if not sample.strip():
        raise OpsError("empty file")
    try:
        dialect = csv.Sniffer().sniff(sample[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(sample), dialect=dialect)
    if not reader.fieldnames:
        raise OpsError("no headers")
    headers = [h.strip() for h in reader.fieldnames if h and h.strip()]
    rows = []
    for rec in reader:
        if not rec:
            continue
        cleaned, empty = {}, True
        for key, value in rec.items():
            if key is None:
                continue
            name = key.strip()
            if isinstance(value, str):
                value = value.strip()
            if value not in (None, ""):
                empty = False
            cleaned[name] = value
        if not empty:
            rows.append(cleaned)
    return headers, rows


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        grid = [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not grid:
        raise OpsError("empty spreadsheet")
    headers = ["" if cell is None else str(cell).strip() for cell in grid[0]]
    rows = []
    for raw in grid[1:]:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        rec = {}
        for i, header in enumerate(headers):
            if not header or i >= len(raw):
                continue
            rec[header] = raw[i]
        rows.append(rec)
    if not any(headers):
        raise OpsError("no headers")
    return [h for h in headers if h], rows


def _lookup(rec, header):
    if header in rec:
        return rec[header]
    want = norm_header(header)
    for key, value in rec.items():
        if norm_header(key) == want:
            return value
    return None


def to_devices(rows, mapping, detected):
    devices = []
    for rec in rows:
        kw = {"source": "CSV import"}
        if detected in _MANAGED:
            kw["managed_by"] = _MANAGED[detected]
        for field, header in (mapping or {}).items():
            if field not in DEVICE_FIELDS:
                continue
            val = _lookup(rec, header)
            if val not in (None, ""):
                kw[field] = val
        if detected == "ad" and kw.get("last_checkin") is not None and str(kw["last_checkin"]).strip().isdigit():
            from connectors.ad import filetime_to_iso
            kw["last_checkin"] = filetime_to_iso(str(kw["last_checkin"]).strip())
        devices.append(device(**kw))
    return devices


def load_import(path, mapping_override=None):
    if not path or not os.path.isfile(path):
        raise OpsError("import file not found")
    headers, rows = read_table(path)
    detected, mapping = detect(headers)
    if mapping_override is not None:
        if not isinstance(mapping_override, dict) or not mapping_override:
            raise OpsError("mapping must be an object")
        mapping = {k: v for k, v in mapping_override.items() if k in DEVICE_FIELDS and isinstance(v, str)}
        if not mapping:
            raise OpsError("mapping is empty")
    return {
        "detected": detected,
        "mapping": mapping,
        "headers": headers,
        "devices": to_devices(rows, mapping, detected),
        "label": _LABELS.get(detected, "CSV import"),
    }
